#!/usr/bin/env python3
"""
可视化并发测试 Web 服务：
- 提供前端页面
- 通过 WebSocket 实时推送每个用户的请求/响应/状态
- 提供启动/停止测试的接口
"""

import asyncio
import hashlib
import io
import json
import re
import secrets
import threading
import time
from pathlib import Path
from typing import Optional, Set, List, Dict

import httpx
from fastapi import FastAPI, UploadFile, File, WebSocket, WebSocketDisconnect, Request
from fastapi.responses import HTMLResponse, JSONResponse

from engine import TestEngine, DEFAULT_CONFIG, parse_accounts, to_password_hash
from record import Recorder
from ui_test import UITestRunner, parse_excel_cases

app = FastAPI(title="测试平台")

# ---- 登录认证 ----
# 默认账号：master / Tuixiang2026
AUTH_USERS = {
    "master": "Tuixiang2026",
}
# 已登录的 token 集合（内存中，重启后失效）
auth_tokens: Set[str] = set()
TOKEN_TTL = 24 * 3600  # token 有效期 24 小时
_token_created: Dict[str, float] = {}

# 全局状态
engine: Optional[TestEngine] = None
engine_task: Optional[asyncio.Task] = None
clients: Set[WebSocket] = set()
event_log: List[dict] = []          # 保留最近事件用于回放
MAX_LOG = 5000
current_test_name: str = ""         # 当前并发测试的名称

# 录制状态
recorder: Optional[Recorder] = None
recorder_task: Optional[asyncio.Task] = None
recorded_requests: List[dict] = []

# UI 自动化测试状态
ui_runner: Optional[UITestRunner] = None
ui_task: Optional[asyncio.Task] = None
ui_cases: List[dict] = []

CONFIG_FILE = Path(__file__).parent / "config.json"
SAVED_FILE = Path(__file__).parent / "saved_tests.json"


def load_config() -> dict:
    """从 config.json 读取配置，不存在则用默认值。"""
    if CONFIG_FILE.exists():
        try:
            saved = json.loads(CONFIG_FILE.read_text(encoding="utf-8"))
            cfg = {**DEFAULT_CONFIG, **saved}
            return cfg
        except Exception:
            pass
    return dict(DEFAULT_CONFIG)


def save_config(cfg: dict) -> None:
    CONFIG_FILE.write_text(json.dumps(cfg, ensure_ascii=False, indent=2),
                           encoding="utf-8")


def load_saved_tests() -> dict:
    """从 saved_tests.json 读取各模块保存的测试配置。"""
    if SAVED_FILE.exists():
        try:
            return json.loads(SAVED_FILE.read_text(encoding="utf-8"))
        except Exception:
            pass
    return {}


def save_saved_tests(data: dict) -> None:
    SAVED_FILE.write_text(json.dumps(data, ensure_ascii=False, indent=2),
                          encoding="utf-8")


def broadcast(event: dict) -> None:
    """把事件推送给所有 WebSocket 客户端。"""
    event_log.append(event)
    if len(event_log) > MAX_LOG:
        del event_log[: len(event_log) - MAX_LOG]
    dead = []
    for ws in clients:
        try:
            asyncio.create_task(ws.send_text(json.dumps(event, ensure_ascii=False)))
        except Exception:
            dead.append(ws)
    for ws in dead:
        clients.discard(ws)


# ---------------- 登录认证 ----------------

def _check_token(token: str) -> bool:
    """校验 token 是否有效。"""
    if not token or token not in auth_tokens:
        return False
    created = _token_created.get(token, 0)
    if time.time() - created > TOKEN_TTL:
        auth_tokens.discard(token)
        _token_created.pop(token, None)
        return False
    return True


@app.middleware("http")
async def auth_middleware(request: Request, call_next):
    """认证中间件：除登录接口和静态资源外，所有请求需携带有效 token。"""
    path = request.url.path
    # 放行登录接口和根路径（根路径返回登录页）
    if path in ("/api/login", "/api/logout", "/") or path.startswith("/static"):
        return await call_next(request)
    # 检查 token（支持 header 或 query 参数）
    token = request.headers.get("Authorization", "").replace("Bearer ", "").strip()
    if not token:
        token = request.query_params.get("token", "")
    if not _check_token(token):
        return JSONResponse({"ok": False, "msg": "未登录或登录已过期"},
                            status_code=401)
    return await call_next(request)


@app.post("/api/login")
async def login(payload: dict) -> dict:
    """登录接口。payload: {username, password}"""
    username = (payload.get("username") or "").strip()
    password = payload.get("password") or ""
    if username not in AUTH_USERS or AUTH_USERS[username] != password:
        return {"ok": False, "msg": "用户名或密码错误"}
    token = secrets.token_hex(32)
    auth_tokens.add(token)
    _token_created[token] = time.time()
    return {"ok": True, "msg": "登录成功", "token": token, "username": username}


@app.post("/api/logout")
async def logout(request: Request) -> dict:
    """登出接口。"""
    token = request.headers.get("Authorization", "").replace("Bearer ", "").strip()
    if not token:
        token = request.query_params.get("token", "")
    auth_tokens.discard(token)
    _token_created.pop(token, None)
    return {"ok": True, "msg": "已退出登录"}


@app.get("/", response_class=HTMLResponse)
async def index() -> str:
    html = Path(__file__).parent / "index.html"
    return html.read_text(encoding="utf-8")


@app.post("/api/proxy")
async def proxy(payload: dict) -> dict:
    """单接口测试代理：转发任意请求，返回完整的请求/响应信息。
    支持 repeat 循环多次，统计平均/最大/最小耗时。
    支持 content_type 指定请求体类型（json/form/text/xml 等）。"""
    url = payload.get("url", "")
    method = (payload.get("method") or "GET").upper()
    body = payload.get("body", "")
    headers = payload.get("headers") or {}
    content_type = payload.get("content_type") or "application/json"
    repeat = int(payload.get("repeat", 1) or 1)
    if repeat < 1:
        repeat = 1
    if not url:
        return {"ok": False, "msg": "缺少 url"}
    import time
    from urllib.parse import urlparse
    parsed = urlparse(url)
    times = []
    last_resp = None
    last_body = ""
    last_ct = ""
    last_size = 0
    last_status = 0
    for _ in range(repeat):
        t0 = time.perf_counter()
        try:
            async with httpx.AsyncClient(timeout=120, follow_redirects=True) as client:
                if method == "POST":
                    if content_type == "application/x-www-form-urlencoded":
                        resp = await client.post(url, data=body,
                                                 headers={"Content-Type": content_type, **headers})
                    else:
                        resp = await client.post(url, content=body,
                                                 headers={"Content-Type": content_type, **headers})
                else:
                    resp = await client.get(url, headers=headers)
            ms = (time.perf_counter() - t0) * 1000
            times.append(round(ms, 1))
            last_resp = resp
            last_status = resp.status_code
            last_ct = resp.headers.get("content-type", "")
            last_size = len(resp.content)
            if "json" in last_ct:
                try:
                    last_body = resp.json()
                except Exception:
                    last_body = resp.text[:5000]
            elif "image" in last_ct or "octet-stream" in last_ct:
                last_body = f"<二进制数据 {len(resp.content)} 字节>"
            else:
                last_body = resp.text[:5000]
        except Exception as e:
            ms = (time.perf_counter() - t0) * 1000
            times.append(round(ms, 1))
            last_status = 0
            last_body = f"{type(e).__name__}: {e}"
    if not last_resp:
        return {"ok": False, "error": last_body,
                "elapsed_ms": round(sum(times) / len(times), 1) if times else 0}
    return {
        "ok": True,
        "request": {
            "method": method,
            "url": url,
            "host": parsed.hostname or "",
            "port": parsed.port or (443 if parsed.scheme == "https" else 80),
            "path": parsed.path or "/",
            "query": parsed.query or "",
            "headers": {"Content-Type": content_type, **headers},
            "body": body,
        },
        "response": {
            "status": last_status,
            "elapsed_ms": round(sum(times) / len(times), 1) if times else 0,
            "avg_ms": round(sum(times) / len(times), 1) if times else 0,
            "max_ms": round(max(times), 1) if times else 0,
            "min_ms": round(min(times), 1) if times else 0,
            "repeat": repeat,
            "times": times,
            "content_type": last_ct,
            "size": last_size,
            "headers": dict(last_resp.headers),
            "body": last_body,
        },
    }


@app.post("/api/multi-test")
async def multi_test(payload: dict) -> dict:
    """多接口测试：按用户输入的接口列表顺序调用。
    每个接口可重复多次，输出完整的请求/响应信息，并统计平均/最大/最小耗时。

    payload:
      base_url: 产品地址
      repeat: 每个接口循环次数
      apis: [{"name": "登录", "method": "POST", "url": "/api/v1/user/login", "body": "{...}"}, ...]
             url 支持占位符 {study_uid} {series_uid} {account} {password_hash}
      study_uid / series_uid / account / password: 占位符替换值
    """
    base = (payload.get("base_url") or "").rstrip("/")
    study_uid = payload.get("study_uid", "")
    series_uid = payload.get("series_uid", "")
    account = payload.get("account", "test")
    password = payload.get("password", "123qwe")
    repeat = int(payload.get("repeat", 1) or 1)
    if repeat < 1:
        repeat = 1
    if not base:
        return {"ok": False, "msg": "缺少 base_url"}

    import time
    from urllib.parse import urlparse
    results = []
    pwd_hash = to_password_hash(password)

    # 用户传入的接口列表；未传则用默认 5 个接口
    apis = payload.get("apis")
    if not apis:
        apis = [
            {"name": "登录", "method": "POST", "url": "/api/v1/user/login",
             "body": '{"name":"{account}","password":"{password_hash}"}'},
            {"name": "图像元数据", "method": "GET",
             "url": "/api/v1/studies?studyInstanceUID={study_uid}&taskType=xa_brain&product=XA_BRAIN"},
            {"name": "序列数据", "method": "GET",
             "url": "/api/repacs/series/{series_uid}/dcp"},
            {"name": "缩略图", "method": "GET",
             "url": "/RESULT/thumbnail/{series_uid}/thumbnail.jpg"},
            {"name": "图像帧", "method": "GET",
             "url": "/xa_brain_encrypt/XA-00001/{study_uid}/{series_uid}/{sop_uid}"},
        ]

    def fill_placeholders(text: str) -> str:
        return (text
                .replace("{study_uid}", study_uid)
                .replace("{series_uid}", series_uid)
                .replace("{account}", account)
                .replace("{password_hash}", pwd_hash))

    async def call(name, method, url, body=None, content_type="application/json"):
        parsed = urlparse(url)
        req_info = {
            "method": method,
            "url": url,
            "host": parsed.hostname or "",
            "port": parsed.port or (443 if parsed.scheme == "https" else 80),
            "path": parsed.path or "/",
            "query": parsed.query or "",
            "headers": {"Content-Type": content_type} if body else {},
            "body": body or "",
        }
        times = []
        last_resp = None
        last_body = ""
        last_ct = ""
        last_size = 0
        last_status = 0
        for _ in range(repeat):
            t0 = time.perf_counter()
            try:
                async with httpx.AsyncClient(timeout=120, follow_redirects=True) as client:
                    if method == "POST":
                        if content_type == "application/x-www-form-urlencoded":
                            resp = await client.post(url, data=body,
                                                     headers={"Content-Type": content_type})
                        else:
                            resp = await client.post(url, content=body,
                                                     headers={"Content-Type": content_type})
                    else:
                        resp = await client.get(url)
                ms = (time.perf_counter() - t0) * 1000
                times.append(round(ms, 1))
                last_resp = resp
                last_status = resp.status_code
                last_ct = resp.headers.get("content-type", "")
                last_size = len(resp.content)
                if "json" in last_ct:
                    try:
                        last_body = resp.json()
                    except Exception:
                        last_body = resp.text[:2000]
                elif "image" in last_ct or "octet-stream" in last_ct:
                    last_body = f"<二进制数据 {len(resp.content)} 字节>"
                else:
                    last_body = resp.text[:2000]
            except Exception as e:
                ms = (time.perf_counter() - t0) * 1000
                times.append(round(ms, 1))
                last_status = 0
                last_body = f"{type(e).__name__}: {e}"
        ok = last_status < 400
        results.append({
            "接口": name,
            "结果": "✅ 通过" if ok else "❌ 失败",
            "请求": req_info,
            "响应": {
                "状态码": last_status,
                "耗时(ms)": round(sum(times) / len(times), 1) if times else 0,
                "平均耗时(ms)": round(sum(times) / len(times), 1) if times else 0,
                "最大耗时(ms)": round(max(times), 1) if times else 0,
                "最小耗时(ms)": round(min(times), 1) if times else 0,
                "循环次数": repeat,
                "每次耗时(ms)": times,
                "大小": last_size,
                "Content-Type": last_ct,
                "响应头": dict(last_resp.headers) if last_resp else {},
                "响应体": last_body,
            },
        })
        return last_resp

    for api in apis:
        name = api.get("name") or api.get("url", "接口")
        method = (api.get("method") or "GET").upper()
        url = api.get("url", "")
        body = api.get("body", "")
        content_type = api.get("content_type") or "application/json"
        if not url:
            results.append({"接口": name, "结果": "⚠️ 跳过",
                            "请求": {"method": method, "url": "", "host": "", "port": 0,
                                     "path": "", "query": "", "headers": {}, "body": ""},
                            "响应": {"状态码": 0, "耗时(ms)": 0, "平均耗时(ms)": 0,
                                     "最大耗时(ms)": 0, "最小耗时(ms)": 0,
                                     "循环次数": repeat, "响应体": "缺少 URL"}})
            continue
        # 相对路径拼上 base_url
        if url.startswith("http://") or url.startswith("https://"):
            full_url = url
        else:
            full_url = base + (url if url.startswith("/") else "/" + url)
        full_url = fill_placeholders(full_url)
        body = fill_placeholders(body) if body else ""
        await call(name, method, full_url, body, content_type)

    passed = sum(1 for r in results if r["结果"].startswith("✅"))
    return {"ok": True, "total": len(results), "passed": passed,
            "repeat": repeat, "results": results}


@app.post("/api/auto-test")
async def auto_test(payload: dict) -> dict:
    """自动化测试：按测试用例步骤顺序执行，支持变量提取和断言。

    payload:
      base_url: 产品地址
      account / password: 登录账号密码
      steps: [{"name": "登录", "method": "POST", "url": "/api/v1/user/login",
               "body": "{...}", "content_type": "application/json",
               "extract": {"token": "data.userId"}, "assert": {"code": 10000}}, ...]

    步骤字段：
      name: 步骤名
      method: GET/POST
      url: 接口路径（支持 {account} {password_hash} {study_uid} {series_uid} 和已提取变量 {var}）
      body: 请求体（同样支持占位符）
      content_type: 请求体类型
      extract: 从响应 JSON 提取变量，如 {"token": "data.userId"}（点号路径）
      assert: 断言，如 {"code": 10000} 或 {"status": 200}
    """
    base = (payload.get("base_url") or "").rstrip("/")
    account = payload.get("account", "test")
    password = payload.get("password", "123qwe")
    study_uid = payload.get("study_uid", "")
    series_uid = payload.get("series_uid", "")
    steps = payload.get("steps") or []
    if not base:
        return {"ok": False, "msg": "缺少 base_url"}
    if not steps:
        return {"ok": False, "msg": "缺少测试步骤"}

    import time
    pwd_hash = to_password_hash(password)
    variables = {
        "account": account,
        "password_hash": pwd_hash,
        "study_uid": study_uid,
        "series_uid": series_uid,
    }

    def fill(text: str) -> str:
        if not text:
            return ""
        for k, v in variables.items():
            text = text.replace("{" + k + "}", str(v))
        return text

    def get_path(obj, path: str):
        """按点号路径从 JSON 取值，如 data.userId。"""
        cur = obj
        for part in path.split("."):
            if isinstance(cur, dict):
                cur = cur.get(part)
            elif isinstance(cur, list) and part.isdigit():
                cur = cur[int(part)]
            else:
                return None
        return cur

    results = []
    async with httpx.AsyncClient(timeout=120, follow_redirects=True) as client:
        for i, step in enumerate(steps):
            name = step.get("name") or f"步骤{i + 1}"
            method = (step.get("method") or "GET").upper()
            url = fill(step.get("url", ""))
            body = fill(step.get("body", ""))
            content_type = step.get("content_type") or "application/json"
            extract = step.get("extract") or {}
            asserts = step.get("assert") or {}

            if url.startswith("/"):
                url = base + url

            t0 = time.perf_counter()
            step_result = {
                "步骤": name,
                "方法": method,
                "URL": url,
                "状态码": 0,
                "耗时(ms)": 0,
                "结果": "❌ 失败",
                "断言": [],
                "提取": {},
                "响应体": "",
            }
            try:
                if method == "POST":
                    if content_type == "application/x-www-form-urlencoded":
                        resp = await client.post(url, data=body,
                                                 headers={"Content-Type": content_type})
                    else:
                        resp = await client.post(url, content=body,
                                                 headers={"Content-Type": content_type})
                else:
                    resp = await client.get(url)
                elapsed = round((time.perf_counter() - t0) * 1000, 1)
                step_result["耗时(ms)"] = elapsed
                step_result["状态码"] = resp.status_code

                # 解析响应体
                ct = resp.headers.get("content-type", "")
                resp_json = None
                if "json" in ct:
                    try:
                        resp_json = resp.json()
                        step_result["响应体"] = resp_json
                    except Exception:
                        step_result["响应体"] = resp.text[:2000]
                elif "image" in ct or "octet-stream" in ct:
                    step_result["响应体"] = f"<二进制数据 {len(resp.content)} 字节>"
                else:
                    step_result["响应体"] = resp.text[:2000]

                # 断言
                all_pass = True
                for key, expected in asserts.items():
                    if key == "status":
                        actual = resp.status_code
                    elif resp_json is not None:
                        actual = get_path(resp_json, key)
                    else:
                        actual = None
                    ok = (actual == expected)
                    if not ok:
                        all_pass = False
                    step_result["断言"].append({
                        "字段": key, "期望": expected, "实际": actual,
                        "结果": "✅" if ok else "❌"})

                # 提取变量
                if resp_json is not None:
                    for var, path in extract.items():
                        val = get_path(resp_json, path)
                        if val is not None:
                            variables[var] = val
                            step_result["提取"][var] = val

                step_result["结果"] = "✅ 通过" if all_pass else "❌ 失败"
            except Exception as e:
                elapsed = round((time.perf_counter() - t0) * 1000, 1)
                step_result["耗时(ms)"] = elapsed
                step_result["响应体"] = f"{type(e).__name__}: {e}"
                step_result["结果"] = "❌ 失败"

            results.append(step_result)

    passed = sum(1 for r in results if r["结果"] == "✅ 通过")
    return {"ok": True, "total": len(results), "passed": passed,
            "results": results, "variables": variables}


@app.post("/api/auto-test/upload")
async def auto_test_upload(file: UploadFile = File(...)) -> dict:
    """上传 Excel 测试用例，自动识别文字步骤，返回解析出的步骤列表。

    支持两种格式：
    1. 每行一个步骤，列顺序：名称|方法|URL|请求体|提取|断言
    2. 任意表格，自动识别含「方法/URL/接口/地址」等关键字的列
    """
    import openpyxl
    content = await file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    except Exception as e:
        return {"ok": False, "msg": f"无法解析 Excel: {type(e).__name__}: {e}"}

    steps = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        # 识别表头
        header = [str(c).strip() if c is not None else "" for c in rows[0]]
        # 尝试匹配列
        col_map = {}
        for idx, h in enumerate(header):
            hl = h.lower()
            if any(k in hl for k in ("名称", "步骤", "name", "用例")):
                col_map.setdefault("name", idx)
            elif any(k in hl for k in ("方法", "method", "请求方式")):
                col_map.setdefault("method", idx)
            elif any(k in hl for k in ("url", "地址", "接口", "路径", "path")):
                col_map.setdefault("url", idx)
            elif any(k in hl for k in ("请求体", "body", "参数", "入参")):
                col_map.setdefault("body", idx)
            elif any(k in hl for k in ("提取", "extract", "变量")):
                col_map.setdefault("extract", idx)
            elif any(k in hl for k in ("断言", "assert", "预期", "期望")):
                col_map.setdefault("assert", idx)

        # 有表头则按列解析，否则按固定顺序
        for row in rows[1:]:
            if row is None or all(c is None or str(c).strip() == "" for c in row):
                continue
            cells = [str(c).strip() if c is not None else "" for c in row]
            if col_map:
                name = cells[col_map["name"]] if "name" in col_map else ""
                method = cells[col_map["method"]] if "method" in col_map else "GET"
                url = cells[col_map["url"]] if "url" in col_map else ""
                body = cells[col_map["body"]] if "body" in col_map else ""
                extract = cells[col_map["extract"]] if "extract" in col_map else ""
                assert_ = cells[col_map["assert"]] if "assert" in col_map else ""
            else:
                name = cells[0] if len(cells) > 0 else ""
                method = cells[1] if len(cells) > 1 else "GET"
                url = cells[2] if len(cells) > 2 else ""
                body = cells[3] if len(cells) > 3 else ""
                extract = cells[4] if len(cells) > 4 else ""
                assert_ = cells[5] if len(cells) > 5 else ""
            if not url:
                continue
            steps.append({
                "name": name or url,
                "method": (method or "GET").upper(),
                "url": url,
                "body": body,
                "extract": extract,
                "assert": assert_,
            })

    if not steps:
        return {"ok": False, "msg": "未识别到测试步骤，请检查 Excel 格式"}

    # 转成文本格式，方便前端填入
    lines = []
    for s in steps:
        lines.append("|".join([s["name"], s["method"], s["url"],
                               s["body"], s["extract"], s["assert"]]))
    return {"ok": True, "count": len(steps), "steps": steps,
            "text": "\n".join(lines)}


@app.websocket("/ws")
async def ws_endpoint(ws: WebSocket) -> None:
    # 校验 token（query 参数）
    token = ws.query_params.get("token", "")
    if not _check_token(token):
        await ws.close(code=4401)
        return
    await ws.accept()
    clients.add(ws)
    # 回放历史事件
    for ev in event_log:
        await ws.send_text(json.dumps(ev, ensure_ascii=False))
    try:
        while True:
            await ws.receive_text()
    except WebSocketDisconnect:
        clients.discard(ws)


@app.get("/api/config")
async def get_config() -> dict:
    cfg = load_config()
    return {"ok": True, "config": cfg,
            "accounts_count": len(cfg.get("accounts", []))}


@app.post("/api/config")
async def set_config(payload: dict) -> dict:
    """保存配置。payload 可含 base_url/password/study_uid/series_uid/product/
    各接口路径/accounts(列表) 或 accounts_text(文本)。"""
    cfg = load_config()
    for key in ("base_url", "password", "study_uid", "series_uid", "product",
                "login_api_path", "check_login_prefix", "studies_api_path",
                "dcp_api_path_template", "frame_prefix", "thumbnail_prefix",
                "msg_token_prefix"):
        if key in payload and payload[key] is not None:
            cfg[key] = str(payload[key]).strip()

    if "accounts" in payload and payload["accounts"]:
        cfg["accounts"] = [str(a).strip() for a in payload["accounts"]
                           if str(a).strip()]
    elif "accounts_text" in payload and payload["accounts_text"]:
        cfg["accounts"] = parse_accounts(payload["accounts_text"])

    save_config(cfg)
    return {"ok": True, "config": cfg,
            "accounts_count": len(cfg.get("accounts", []))}


# ---------------- 测试配置保存/加载 ----------------

@app.get("/api/saved-tests")
async def get_saved_tests() -> dict:
    """返回各模块保存的测试配置历史列表。"""
    return {"ok": True, "saved": load_saved_tests()}


@app.post("/api/saved-tests")
async def set_saved_tests(payload: dict) -> dict:
    """保存某个模块的一条测试配置（追加到历史记录）。

    payload:
      module: single / multi / concurrency / auto
      name:   记录名称（可选，默认用时间）
      data:   该模块的表单数据（任意 JSON）
    """
    module = (payload.get("module") or "").strip()
    if module not in ("single", "multi", "concurrency", "auto"):
        return {"ok": False, "msg": "未知模块"}
    data = payload.get("data")
    if data is None:
        return {"ok": False, "msg": "缺少 data"}
    import time as _time
    name = (payload.get("name") or "").strip()
    if not name:
        name = _time.strftime("%m-%d %H:%M:%S")
    record = {
        "id": str(int(_time.time() * 1000)),
        "name": name,
        "time": _time.strftime("%Y-%m-%d %H:%M:%S"),
        "data": data,
    }
    saved = load_saved_tests()
    # 兼容旧格式：旧数据是 dict，转成列表
    old = saved.get(module)
    if isinstance(old, dict):
        old = [{"id": "legacy", "name": "旧记录", "time": "", "data": old}]
    elif isinstance(old, list):
        old = old
    else:
        old = []
    old.insert(0, record)
    saved[module] = old
    save_saved_tests(saved)
    return {"ok": True, "msg": "已保存", "record": record, "saved": saved}


@app.delete("/api/saved-tests/{module}/{record_id}")
async def delete_saved_test(module: str, record_id: str) -> dict:
    """删除某个模块的一条保存记录。"""
    saved = load_saved_tests()
    records = saved.get(module)
    if isinstance(records, list):
        new_records = [r for r in records if r.get("id") != record_id]
        if len(new_records) != len(records):
            saved[module] = new_records
            save_saved_tests(saved)
            return {"ok": True, "msg": "已删除"}
    return {"ok": False, "msg": "记录不存在"}


@app.post("/api/start")
async def start_test(payload: dict) -> dict:
    global engine, engine_task, current_test_name
    if engine_task and not engine_task.done():
        return {"ok": False, "msg": "测试已在运行中"}

    users = int(payload.get("users", 20))
    observe = float(payload.get("observe", 300))
    mode = payload.get("mode", "full")
    current_test_name = (payload.get("name") or "").strip()
    event_log.clear()

    # 优先用请求里带的配置，否则用已保存配置
    cfg = load_config()
    if "config" in payload and payload["config"]:
        cfg = {**cfg, **payload["config"]}
        if "accounts_text" in payload["config"]:
            cfg["accounts"] = parse_accounts(payload["config"]["accounts_text"])
    # 录制导入的自定义接口列表
    if payload.get("apis"):
        cfg["apis"] = payload["apis"]

    engine = TestEngine(users=users, observe_seconds=observe,
                        emit=broadcast, config=cfg, mode=mode)
    engine_task = asyncio.create_task(engine.run())
    label = "图像帧接口" if mode == "frame_only" else ("自定义接口" if payload.get("apis") else "全部接口")
    return {"ok": True, "msg": f"已启动 {users} 用户并发测试（{label}）"}


@app.post("/api/stop")
async def stop_test() -> dict:
    global engine
    if engine:
        engine.stop()
    return {"ok": True, "msg": "已请求停止"}


@app.get("/api/status")
async def status() -> dict:
    running = bool(engine_task and not engine_task.done())
    return {"running": running, "events": len(event_log)}


@app.post("/api/clear")
async def clear_log() -> dict:
    """清空历史事件日志。"""
    event_log.clear()
    return {"ok": True, "msg": "日志已清空"}


# ---------------- 录制 ----------------

@app.post("/api/record/start")
async def record_start(payload: dict) -> dict:
    """启动录制：打开真实浏览器，用户手动操作，捕获请求。
    支持 login_first：先登录注入 cookie，再打开目标地址。"""
    global recorder, recorder_task, recorded_requests
    if recorder_task and not recorder_task.done():
        return {"ok": False, "msg": "录制已在进行中"}
    url = (payload.get("url") or "").strip()
    if not url:
        return {"ok": False, "msg": "缺少 url"}
    login_first = bool(payload.get("login_first"))
    cfg = load_config()
    base_url = (payload.get("base_url") or cfg.get("base_url") or "").strip()
    account = (payload.get("account") or "").strip() or "test"
    password = payload.get("password") or cfg.get("password") or "123qwe"
    login_api_path = cfg.get("login_api_path", "/api/v1/user/login")
    recorded_requests = []
    recorder = Recorder(emit=broadcast)
    recorder_task = asyncio.create_task(
        recorder.run(url, login_first=login_first, base_url=base_url,
                     account=account, password=password,
                     login_api_path=login_api_path))
    return {"ok": True, "msg": "录制已启动，请在浏览器中操作"}


@app.post("/api/record/stop")
async def record_stop() -> dict:
    """停止录制，返回捕获的请求列表。"""
    global recorder, recorder_task, recorded_requests
    if not recorder:
        return {"ok": False, "msg": "尚未开始录制"}
    recorder.stop()
    if recorder_task:
        try:
            recorded_requests = await asyncio.wait_for(recorder_task, timeout=30)
        except Exception:
            recorded_requests = recorder.requests
    recorder = None
    recorder_task = None
    return {"ok": True, "count": len(recorded_requests),
            "requests": recorded_requests}


@app.get("/api/record/status")
async def record_status() -> dict:
    running = bool(recorder_task and not recorder_task.done())
    return {"running": running, "count": len(recorded_requests)}


@app.post("/api/record/import")
async def record_import(payload: dict) -> dict:
    """接收本地录制脚本上传的请求列表。"""
    global recorded_requests
    requests = payload.get("requests") or []
    if not requests:
        return {"ok": False, "msg": "没有请求数据"}
    recorded_requests = requests
    # 广播给前端，让录制页面实时显示
    for r in requests:
        broadcast({"type": "record_request", "method": r.get("method", ""),
                   "url": r.get("url", ""), "ts": time.time()})
    return {"ok": True, "count": len(requests), "requests": requests}


@app.get("/api/record/imported")
async def record_imported() -> dict:
    """返回最近一次导入的请求列表。"""
    return {"ok": True, "requests": recorded_requests}


@app.post("/api/record/import-har")
async def record_import_har(file: UploadFile = File(...)) -> dict:
    """接收 Chrome 开发者工具导出的 HAR 文件，解析出请求列表。"""
    global recorded_requests
    content = await file.read()
    try:
        har = json.loads(content.decode("utf-8"))
    except Exception as e:
        return {"ok": False, "msg": f"无法解析 HAR 文件: {type(e).__name__}: {e}"}

    entries = har.get("log", {}).get("entries", [])
    requests = []
    for e in entries:
        req = e.get("request", {})
        url = req.get("url", "")
        # 跳过静态资源
        if re.search(r'\.(js|css|png|jpg|jpeg|gif|svg|ico|woff|woff2|ttf|map)(\?|$)', url, re.I):
            continue
        method = req.get("method", "GET")
        headers = {}
        for h in req.get("headers", []):
            headers[h.get("name", "")] = h.get("value", "")
        body = ""
        post_data = req.get("postData", {})
        if post_data:
            body = post_data.get("text", "")
        requests.append({
            "method": method,
            "url": url,
            "headers": headers,
            "body": body,
            "ts": e.get("startedDateTime", ""),
        })

    if not requests:
        return {"ok": False, "msg": "HAR 文件中没有可导入的 API 请求"}

    recorded_requests = requests
    for r in requests:
        broadcast({"type": "record_request", "method": r.get("method", ""),
                   "url": r.get("url", ""), "ts": time.time()})
    return {"ok": True, "count": len(requests), "requests": requests}


# ---------------- UI 自动化测试 ----------------

@app.post("/api/ui-test/upload")
async def ui_test_upload(file: UploadFile = File(...)) -> dict:
    """上传 Excel 用例，解析并返回用例列表。"""
    global ui_cases
    content = await file.read()
    try:
        cases = parse_excel_cases(content)
    except Exception as e:
        return {"ok": False, "msg": f"无法解析 Excel: {type(e).__name__}: {e}"}
    if not cases:
        return {"ok": False, "msg": "未识别到用例，请检查 Excel 格式（需含：编号/名称/操作步骤/预期结果）"}
    ui_cases = cases
    return {"ok": True, "count": len(cases), "cases": cases}


@app.post("/api/ui-test/start")
async def ui_test_start(payload: dict) -> dict:
    """启动 UI 自动化测试：用浏览器按用例步骤执行。

    payload:
      base_url: 起始地址（可选，若用例第一步不是打开网址则先打开它）
      headless: 是否无头模式（默认 true）
      login_url: 登录页地址（自动登录时打开）
      accounts: {"角色名": {"account": ..., "password": ...}}
      cases: 用例列表（可选，不传则用上次上传的）
    """
    global ui_runner, ui_task, ui_cases
    if ui_task and not ui_task.done():
        return {"ok": False, "msg": "UI 测试已在运行中"}
    base_url = (payload.get("base_url") or "").strip()
    headless = bool(payload.get("headless", True))
    login_url = (payload.get("login_url") or "").strip()
    accounts = payload.get("accounts") or {}
    step_interval = float(payload.get("step_interval", 0.5) or 0.5)
    cases = payload.get("cases") or ui_cases
    if not cases:
        return {"ok": False, "msg": "没有用例，请先上传 Excel"}
    ui_runner = UITestRunner(emit=broadcast)
    ui_task = asyncio.create_task(
        ui_runner.run(cases, base_url, headless, accounts, login_url,
                      step_interval))
    return {"ok": True, "msg": f"已启动 UI 测试，共 {len(cases)} 条用例"}


@app.post("/api/ui-test/stop")
async def ui_test_stop() -> dict:
    """停止 UI 自动化测试。"""
    global ui_runner
    if ui_runner:
        ui_runner.stop()
    return {"ok": True, "msg": "已请求停止"}


@app.get("/api/ui-test/status")
async def ui_test_status() -> dict:
    running = bool(ui_task and not ui_task.done())
    return {"running": running, "cases": len(ui_cases)}


@app.get("/api/ui-test/cases")
async def ui_test_cases() -> dict:
    """返回当前已上传的用例列表。"""
    return {"ok": True, "cases": ui_cases}


@app.get("/api/ui-test/results")
async def ui_test_results() -> dict:
    """返回最近一次 UI 测试的结果。"""
    global ui_runner
    results = []
    if ui_runner is not None:
        # 从事件日志中提取 ui_test_done 事件的结果
        for ev in reversed(event_log):
            if ev.get("type") == "ui_test_done":
                results = ev.get("results", [])
                break
    return {"ok": True, "results": results}


@app.get("/api/report")
async def report() -> dict:
    """生成可审计报告：每个用户的账号、独立会话 token、精确时间戳。"""
    if not engine:
        return {"ok": False, "msg": "尚未运行测试"}
    states = engine._snapshot()
    # 并发度分析：登录时间戳的极差、阅片时间戳的极差
    login_ts = [s["login_ts"] for s in states if s["login_ts"] > 0]
    viewer_ts = [s["viewer_ts"] for s in states if s["viewer_ts"] > 0]
    first_ts = [s["first_frame_ts"] for s in states if s["first_frame_ts"] > 0]

    def spread(vals):
        if not vals:
            return 0.0
        return round(max(vals) - min(vals), 3)

    tokens = [s["session_token"] for s in states if s["session_token"]]
    return {
        "ok": True,
        "users": len(states),
        "mode": getattr(engine, "mode", "full"),
        "unique_tokens": len(set(tokens)),
        "login_spread_seconds": spread(login_ts),
        "viewer_spread_seconds": spread(viewer_ts),
        "first_frame_spread_seconds": spread(first_ts),
        "states": states,
    }


@app.get("/api/aggregate")
async def aggregate() -> dict:
    """聚合报告：按接口统计，字段与单接口测试一致
    （URL/主机/端口/方法/路径/状态码/结果/循环次数/平均耗时/最小耗时/最大耗时/大小）。"""
    import statistics
    from urllib.parse import urlparse
    # 从事件日志中聚合
    reqs = [e for e in event_log if e.get("type") == "request"]
    resps = [e for e in event_log if e.get("type") == "response"]

    # 按 tag 记录 method（从 request 事件里取）
    tag_method: dict = {}
    for r in reqs:
        tag = r.get("tag", "其他")
        if tag not in tag_method:
            tag_method[tag] = r.get("method", "GET")

    # 按 tag 分组统计响应
    by_tag: dict = {}
    for r in resps:
        tag = r.get("tag", "其他")
        if tag not in by_tag:
            by_tag[tag] = {"耗时": [], "大小": [], "状态码": [], "url": "",
                           "成功": 0, "失败": 0}
        info = by_tag[tag]
        status = r.get("status", 0)
        if 0 < status < 400:
            info["成功"] += 1
        else:
            info["失败"] += 1
        info["状态码"].append(status)
        elapsed = r.get("elapsed_ms", 0)
        if elapsed > 0:
            info["耗时"].append(elapsed)
        size = r.get("size", 0)
        if size > 0:
            info["大小"].append(size)
        if not info["url"]:
            info["url"] = r.get("url", "")

    def avg(vals):
        return round(statistics.mean(vals), 1) if vals else 0

    # 接口汇总（字段与单接口测试一致）
    api_summary = []
    for tag, info in by_tag.items():
        parsed = urlparse(info["url"])
        # 多数状态码
        status = max(set(info["状态码"]), key=info["状态码"].count) if info["状态码"] else 0
        ok = 0 < status < 400
        api_summary.append({
            "接口": tag,
            "URL": info["url"],
            "主机": parsed.hostname or "",
            "端口": parsed.port or (443 if parsed.scheme == "https" else 80),
            "方法": tag_method.get(tag, "GET"),
            "路径": parsed.path or "/",
            "状态码": status,
            "结果": "✅ 通过" if ok else "❌ 失败",
            "循环次数": info["成功"] + info["失败"],
            "平均耗时": avg(info["耗时"]),
            "最小耗时": round(min(info["耗时"]), 1) if info["耗时"] else 0,
            "最大耗时": round(max(info["耗时"]), 1) if info["耗时"] else 0,
            "大小": avg(info["大小"]),
        })

    # 状态统计
    states = engine._snapshot() if engine else []
    return {
        "ok": True,
        "name": current_test_name,
        "mode": getattr(engine, "mode", "full") if engine else "full",
        "users": len(states),
        "总耗时(秒)": round(max([e.get("ts", 0) for e in resps if e.get("ts")]) -
                             min([e.get("ts", 0) for e in reqs if e.get("ts")]), 1)
                            if reqs and resps else 0.0,
        "接口汇总": api_summary,
        "状态统计": {
            "完成": sum(1 for s in states if s["status"] == "done"),
            "超时": sum(1 for s in states if s["status"] == "timeout"),
            "出错": sum(1 for s in states if s["status"] == "error"),
            "加载中": sum(1 for s in states if s["status"] == "loading"),
        },
    }


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=19000, log_level="warning")
